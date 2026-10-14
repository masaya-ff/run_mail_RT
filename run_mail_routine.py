"""メール受信ルーティン.xlsx に記載されたフォルダの bat ファイルを順に実行する。

B列（該当）に「〇」が付いている行のみ対象。D列のフォルダ内にある *.bat を実行する。
"""
import subprocess
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

BASE_DIR = Path(__file__).resolve().parent
EXCEL_PATH = BASE_DIR / "メール受信ルーティン.xlsx"
MARKS = {"〇", "○", "◯"}  # 表記ゆれを許容
FIRST_DATA_ROW = 3  # 1行目:タイトル, 2行目:見出し


def load_targets():
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    targets = []
    for row_no, row in enumerate(
        ws.iter_rows(min_row=FIRST_DATA_ROW, min_col=2, max_col=4, values_only=True),
        start=FIRST_DATA_ROW,
    ):
        mark, name, folder = row
        if str(mark or "").strip() in MARKS and folder:
            targets.append((row_no, str(name or ""), Path(str(folder).strip())))
    wb.close()
    return targets


def write_run_times(run_times):
    """実行した行のE列に実行日時を書き込む。"""
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.worksheets[0]
    if not ws.cell(row=FIRST_DATA_ROW - 1, column=5).value:
        ws.cell(row=FIRST_DATA_ROW - 1, column=5, value="最終実行日時")
    for row_no, run_at in run_times.items():
        cell = ws.cell(row=row_no, column=5, value=run_at)
        cell.number_format = "yyyy/mm/dd hh:mm"
    ws.column_dimensions["E"].width = max(ws.column_dimensions["E"].width or 0, 18)
    while True:
        try:
            wb.save(EXCEL_PATH)
            print("実行日時をExcelのE列に記録しました。")
            return True
        except PermissionError:
            print("※Excelが開かれているため、実行日時を書き込めません。")
            answer = input("  Excelを閉じてEnterを押してください（記録しない場合は n + Enter）: ")
            if answer.strip().lower() == "n":
                print("  実行日時は記録しませんでした。")
                return False


def main():
    targets = load_targets()
    if not targets:
        print("実行対象がありません（B列に〇が付いた行なし）。")
        return 0

    failed = []
    run_times = {}
    for row_no, name, folder in targets:
        print("=" * 60)
        print(f"[{name}] {folder}")
        if not folder.is_dir():
            print("  → フォルダが見つかりません。スキップします。")
            failed.append(name)
            continue
        bats = sorted(folder.glob("*.bat"))
        if not bats:
            print("  → batファイルが見つかりません。スキップします。")
            failed.append(name)
            continue
        run_times[row_no] = datetime.now().replace(second=0, microsecond=0)
        for bat in bats:
            print(f"  実行: {bat.name}")
            # 各batの末尾の pause で止まらないよう、標準入力に改行を渡す
            result = subprocess.run(
                ["cmd", "/c", str(bat)],
                cwd=folder,
                input=b"\r\n",
            )
            if result.returncode != 0:
                print(f"  → 終了コード {result.returncode}")
                failed.append(f"{name} ({bat.name})")

    print("=" * 60)
    if run_times:
        write_run_times(run_times)
    if failed:
        print("エラー／スキップ:")
        for f in failed:
            print(f"  - {f}")
        return 1
    print("すべて完了しました。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
